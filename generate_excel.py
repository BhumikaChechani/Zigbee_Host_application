import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

md_text = """
### 🟢 1. Aqara Wireless Mini Switch (Smart Button)
| Test Action (What to do) | Expected Result (System Response) | Notes |
| :--- | :--- | :--- |
| Press button 1 time | Console logs single press. No siren. | Debounce is set to 0.15s. |
| Press button 2 times | Console logs double press. No siren. | |
| Press button 3 times FAST (within 3 secs) | FULL SIREN ALARM on all sirens. | Simulates an SOS / Panic action. |
| Press button 3 times SLOW (> 3 secs gap) | Console resets history. No alarm. | Tests timeout rejection logic. |

### 🔴 2. Onics Panic Button
| Test Action (What to do) | Expected Result (System Response) | Notes |
| :--- | :--- | :--- |
| Press and hold button (Panic Set) | FULL SIREN ALARM on all sirens. | Immediate emergency trigger. |
| Release button (Panic Clear) | ALL SIRENS STOP immediately. | |
| Type env <addr> in CLI | Prints Battery/Voltage & Temperature. | Tests generic sensor health data. |

### 🚪 3. Contact Sensor (Door/Window)
| Test Action (What to do) | Expected Result (System Response) | Notes |
| :--- | :--- | :--- |
| Move magnet away (Open Door) | 1-SECOND CHIME on sirens. | HARDWARE LIMITATION: Siren ignores low-volume setting and always plays loud. |
| Bring magnet close (Close Door) | Console logs door closed. No chime. | |
| Type env <addr> in CLI | Prints Battery/Voltage & Temperature. | Tests generic sensor health data. |

### 📳 4. Frient Vibration Sensor (Glass Break)
| Test Action (What to do) | Expected Result (System Response) | Notes |
| :--- | :--- | :--- |
| Tap glass gently (Vibration / Alarm 2) | 1-SECOND LOUD BEEP on sirens. | HARDWARE LIMITATION: Siren ignores low-volume setting and always plays loud. |
| Shake or tilt heavily (Movement / Alarm 1) | FULL SIREN ALARM on all sirens. | Critical breach / window broken. |
| Remove battery cover (Tamper switch) | FULL SIREN ALARM on all sirens. | Protects device from being dismantled. |
| Type env <addr> in CLI | Prints Battery/Voltage & Temperature. | WAKEUP REQUIRED: You must tap/vibrate the sensor right before hitting enter to wake its radio! |
| Type sensitivity <addr> <1-15> | Configures hardware sensitivity. | 1=Most sensitive, 15=Least sensitive (Default 10). |

### 🚨 5. Frient Smart Siren (SIRZB-110)
| Test Action (What to do) | Expected Result (System Response) | Notes |
| :--- | :--- | :--- |
| Remove mounting backplate (Tamper switch) | FULL SIREN ALARM on all sirens. | Siren hardware tamper detection. |
| Type env <addr> in CLI | Prints Battery/Voltage. | Fetches diagnostics from the siren. |
| Type siren on and hit Enter | FULL SIREN ALARM on all sirens. | Tests global siren activation. |
| Type siren off and hit Enter | ALL SIRENS STOP immediately. | Stops the global test. |
| Type siren mode <1-6> | Sets global siren sound mode. | Modes: 1=Burglar, 2=Fire, 3=Emergency, 4=Police Panic, 5=Fire Panic, 6=Emergency Panic. |
| Type siren vol <0-3> | Sets global siren volume level. | Levels: 0=Low, 1=Medium, 2=High, 3=Very High. Note: Some firmwares ignore volume. |
| Type siren test <addr> [mode] | Sounds ONLY the specific siren address. | Provide mode 1-6 to test a specific sound. |
| Type siren stop <addr> | ONLY THAT SPECIFIC SIREN stops. | Stops a siren test. |

### 🚶 6. Aqara FP300 Presence Sensor
| Test Action (What to do) | Expected Result (System Response) | Notes |
| :--- | :--- | :--- |
| Walk into the room / zone | Console logs person detected. | Siren logic is disabled for testing. |
| Type env <addr> in CLI | Prints Temperature & Humidity. | Tests FP300 environmental data. |
| Type zone <addr> <idx> <start> <end> | Configures a specific detection zone. | Start/End are in 25cm slices. |
| Type zonedel <addr> <idx> | Deletes the specific detection zone. | |
| Type spatiallearn <addr> | Triggers spatial background learning. | Ensure the room is empty first! |
| Type sensitivity <addr> <1|2|3> | Configures radar sensitivity. | 1=Low, 2=Medium, 3=High. |

### 💻 7. Global CLI Commands (Controller Level)
| Test Action (What to do) | Expected Result (System Response) | Notes |
| :--- | :--- | :--- |
| Type status and hit Enter | Prints list of all connected devices. | Shows online status and last-seen time. |
| Type permit [seconds] | Opens the Zigbee network for pairing. | Default is 60s if not specified. |
| Type discover <addr> | Discovers device endpoints and clusters. | Forces Zigbee active endpoint discovery. |
| Type forcesetup <addr> | Re-runs initial configuration binding. | Fixes devices that didn't set up correctly. |
"""

wb = Workbook()
ws = wb.active
ws.title = "QA Testing Matrix"

title_font = Font(size=14, bold=True, color="FFFFFF")
title_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
header_font = Font(bold=True)
header_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")

current_row = 1
lines = md_text.split("\n")

for line in lines:
    line = line.strip()
    if line.startswith("### "):
        category = line.replace("### ", "").strip()
        cell = ws.cell(row=current_row, column=1, value=category)
        cell.font = title_font
        cell.fill = title_fill
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
        current_row += 1
    elif line.startswith("| Test Action"):
        headers = ["Test Action (What to do)", "Expected Result (System Response)", "Notes"]
        for col_idx, header in enumerate(headers, 1):
            c = ws.cell(row=current_row, column=col_idx, value=header)
            c.font = header_font
            c.fill = header_fill
        current_row += 1
    elif line.startswith("| :---"):
        pass
    elif line.startswith("|"):
        parts = [p.strip() for p in line.split("|")[1:-1]]
        if len(parts) == 3:
            ws.cell(row=current_row, column=1, value=parts[0])
            ws.cell(row=current_row, column=2, value=parts[1])
            ws.cell(row=current_row, column=3, value=parts[2])
            current_row += 1
    elif line == "":
        if current_row > 1 and ws.cell(row=current_row-1, column=1).value != None:
            current_row += 2 # Leave 2 empty rows between tables

for col in range(1, 4):
    max_length = 0
    column_letter = get_column_letter(col)
    for row in range(1, current_row):
        try:
            cell_value = ws.cell(row=row, column=col).value
            if cell_value and not str(cell_value).startswith("🟢") and not str(cell_value).startswith("🔴") and not str(cell_value).startswith("🚪") and not str(cell_value).startswith("📳") and not str(cell_value).startswith("🚨") and not str(cell_value).startswith("🚶") and not str(cell_value).startswith("💻"):
                max_length = max(max_length, len(str(cell_value)))
        except:
            pass
    adjusted_width = min(max_length + 2, 85)
    if adjusted_width < 15: adjusted_width = 35
    ws.column_dimensions[column_letter].width = adjusted_width

wb.save("QA_TEST_MATRIX.xlsx")
print("Successfully generated beautifully formatted single-sheet Excel file matching markdown explicitly!")
